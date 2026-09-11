import json
import logging
import re
from pathlib import Path

import pandas as pd
from PIL import Image, ImageDraw
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


Image.MAX_IMAGE_PIXELS = None

UNCERTAIN_FILL   = PatternFill(fill_type="solid", start_color="FFFF00", end_color="FFFF00")
HEADER_FILL      = PatternFill(fill_type="solid", start_color="1A365D", end_color="1A365D")
ROW_ALT_FILL     = PatternFill(fill_type="solid", start_color="F0F4F8", end_color="F0F4F8")
HEADER_FONT      = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
CELL_FONT        = Font(name="Calibri", size=10)
HEADER_ALIGN     = Alignment(horizontal="center", vertical="center", wrap_text=False)
CELL_ALIGN       = Alignment(horizontal="left",   vertical="center", wrap_text=False)
THIN_BORDER      = Border(
    left=Side(style="thin", color="D9E2EC"),
    right=Side(style="thin", color="D9E2EC"),
    top=Side(style="thin", color="D9E2EC"),
    bottom=Side(style="thin", color="D9E2EC"),
)

COLUMN_WIDTHS = {
    "drop_id": 16, "row": 6, "col": 6, "camera_id": 10,
    "frame_rank": 11, "time_seconds": 13, "time_label": 14,
    "image_id": 12, "image_file": 32,
    "x_image": 10, "y_image": 10, "x_data": 10, "y_data": 10,
    "x_in_cell": 12, "y_in_cell": 12,
    "species_common_name": 26, "species_scientific_name": 30,
    "family": 18, "species_color": 14,
    "count": 8, "comment": 26, "is_uncertain": 13,
    "user": 14, "timestamp": 22,
    "campaign": 16, "location": 16, "site": 14,
    "date": 13, "latitude": 12, "longitude": 12,
    "habitat": 18, "depth": 13, "visibility": 14,
}

ANNOTATION_COLUMNS = [
    # — Metadata —
    "campaign",
    "location",
    "site",
    "latitude",
    "longitude",
    "drop_id",
    "date",
    # — Caméra / Temps —
    "camera_id",
    "frame_rank",
    "time_seconds",
    "time_label",
    # — Spatial (fish.annot) —
    "row",
    "col",
    "image_id",
    "image_file",
    "x_image",
    "y_image",
    "x_data",
    "y_data",
    "x_in_cell",
    "y_in_cell",
    # — Biologie —
    "species_common_name",
    "species_scientific_name",
    "family",
    "species_color",
    "count",
    # — Annotation —
    "comment",
    "is_uncertain",
    "user",
    "timestamp",
    # — Conditions terrain —
    "habitat",
    "visibility",
    "depth",
]


def save_csv_xlsx(annotations, folder):
    folder = Path(folder)
    csv_path = folder / "annotations.csv"
    xlsx_path = folder / "annotations.xlsx"

    if annotations:
        df = pd.DataFrame(annotations)
    else:
        df = pd.DataFrame(columns=ANNOTATION_COLUMNS)

    for col in ANNOTATION_COLUMNS:
        if col not in df.columns:
            df[col] = ""

    df = df[ANNOTATION_COLUMNS]

    df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    df.to_excel(xlsx_path, index=False)

    _format_excel(xlsx_path)


def load_existing_annotations(folder):
    folder = Path(folder)
    csv_path = folder / "annotations.csv"
    if not csv_path.exists():
        return []

    # keep_default_na=False: without it every blank cell comes back as float NaN,
    # which then reaches the annotation table as the literal text "nan" and can be
    # written back into the data when the row is edited.
    df = pd.read_csv(csv_path, keep_default_na=False, dtype=str)

    annotations = []
    for _, row in df.iterrows():
        ann = row.to_dict()

        for key in ["row", "col", "camera_id", "frame_rank", "time_seconds", "count"]:
            if key in ann and pd.notna(ann[key]):
                try:
                    ann[key] = int(float(ann[key]))
                except Exception:
                    pass

        for key in ["x_image", "y_image", "x_data", "y_data", "x_in_cell", "y_in_cell"]:
            if key in ann and pd.notna(ann[key]):
                try:
                    ann[key] = float(ann[key])
                except Exception:
                    pass

        ann["is_uncertain"] = _to_bool(ann.get("is_uncertain", False))
        annotations.append(ann)

    return annotations


def _to_bool(value):
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"true", "1", "yes", "oui"}


def _format_excel(xlsx_path):
    wb = load_workbook(xlsx_path)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    uncertain_col_index = next(
        (i for i, h in enumerate(headers, 1) if h == "is_uncertain"), None
    )

    # header row
    ws.row_dimensions[1].height = 22
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill   = HEADER_FILL
        cell.font   = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = COLUMN_WIDTHS.get(header, 14)

    # data rows
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 16
        is_uncertain = (
            uncertain_col_index is not None
            and _to_bool(ws.cell(row=row_idx, column=uncertain_col_index).value)
        )
        row_fill = UNCERTAIN_FILL if is_uncertain else (
            ROW_ALT_FILL if row_idx % 2 == 0 else None
        )
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font      = CELL_FONT
            cell.alignment = CELL_ALIGN
            cell.border    = THIN_BORDER
            if row_fill:
                cell.fill = row_fill

    # freeze header + auto-filter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(xlsx_path)


MIN_CROP_PX = 80  # taille minimale (px) de la plus petite dimension du crop


def export_crops_par_espece(annotations_list, drop_folder, output_folder):
    """
    Découpe un crop par annotation et organise dans output_folder/{nom_scientifique}/.
    Lit les frames pleine résolution écrites par l'onglet Prepare quand « Save source
    frames » est coché ({cam_id}_frames/frame_NNNNNN.jpg).
    Si un dossier {cam_id}_detections/ fournit des boîtes englobantes (produites par
    un outil externe, voir _trouver_bbox_proche), utilise la plus proche du clic dans
    un rayon de 200 px ; sinon carré de 200×200 px centré sur le clic.
    Filtre les crops dont la plus petite dimension est < MIN_CROP_PX.
    Applique un padding carré (fond noir) avant la sauvegarde.
    Retourne {"n_bbox": int, "n_fallback": int, "n_ignore": int, "n_trop_petit": int}.
    """
    drop_folder = Path(drop_folder)
    output_folder = Path(output_folder)

    cell_w, cell_h = _lire_cell_dims(drop_folder)

    n_bbox = 0
    n_fallback = 0
    n_ignore = 0
    n_trop_petit = 0

    for ann in annotations_list:
        camera_id = ann.get("camera_id")
        frame_rank = ann.get("frame_rank")
        image_id = str(ann.get("image_id", "")).strip()
        species_common = str(ann.get("species_common_name", "")).strip()
        species_scientific = str(ann.get("species_scientific_name", "")).strip()

        try:
            x_in_cell = float(ann.get("x_in_cell"))
            y_in_cell = float(ann.get("y_in_cell"))
            frame_rank = int(frame_rank)
            camera_id = str(camera_id)
        except (TypeError, ValueError):
            logging.warning(f"Annotation ignorée (champs invalides) : image_id={image_id}")
            n_ignore += 1
            continue

        # Dossier de sortie : _HESITATION si hésitation, sinon nom scientifique nettoyé
        if "HESITATION" in species_common.upper():
            nom_dossier = "_HESITATION"
        else:
            nom_dossier = _nettoyer_nom_dossier(species_scientific)

        # Frame source pleine résolution dans {camera_id}_frames/
        frame_path = drop_folder / f"{camera_id}_frames" / f"frame_{frame_rank:06d}.jpg"
        if not frame_path.exists():
            logging.warning(f"Frame introuvable : {frame_path}")
            n_ignore += 1
            continue

        try:
            img = Image.open(frame_path).convert("RGB")
        except Exception as e:
            logging.warning(f"Impossible de charger {frame_path} : {e}")
            n_ignore += 1
            continue

        img_w, img_h = img.size

        # Convertir les coordonnées du clic (espace cellule) → espace frame original
        if cell_w and cell_h:
            x_click = x_in_cell * img_w / cell_w
            y_click = y_in_cell * img_h / cell_h
        else:
            x_click = x_in_cell
            y_click = y_in_cell

        # Recherche de la boîte englobante la plus proche du clic (coordonnées frame)
        detections_folder = drop_folder / f"{camera_id}_detections"
        frame_key = f"frame_{frame_rank:06d}"
        bbox = _trouver_bbox_proche(detections_folder, frame_key, x_click, y_click, seuil_px=200)

        if bbox is not None:
            x1_bb, y1_bb, x2_bb, y2_bb = bbox
            n_bbox += 1
        else:
            # Fallback : carré 200×200 px centré sur le clic (en coordonnées frame)
            x1_bb = x_click - 100
            y1_bb = y_click - 100
            x2_bb = x_click + 100
            y2_bb = y_click + 100
            n_fallback += 1

        # Marge de 20% de chaque côté pour avoir du contexte, clippée aux bords
        larg_bb = x2_bb - x1_bb
        haut_bb = y2_bb - y1_bb
        x1_crop = max(0, int(x1_bb - larg_bb * 0.20))
        y1_crop = max(0, int(y1_bb - haut_bb * 0.20))
        x2_crop = min(img_w, int(x2_bb + larg_bb * 0.20))
        y2_crop = min(img_h, int(y2_bb + haut_bb * 0.20))

        crop_w = x2_crop - x1_crop
        crop_h = y2_crop - y1_crop

        # Filtre taille minimale
        if crop_w < MIN_CROP_PX or crop_h < MIN_CROP_PX:
            logging.warning(
                f"Crop ignoré (trop petit {crop_w}×{crop_h} px) : {image_id}"
            )
            n_trop_petit += 1
            continue

        crop = img.crop((x1_crop, y1_crop, x2_crop, y2_crop))

        # Padding carré (fond noir) pour uniformiser la forme des crops
        side = max(crop_w, crop_h)
        if crop_w != crop_h:
            padded = Image.new("RGB", (side, side), (0, 0, 0))
            padded.paste(crop, ((side - crop_w) // 2, (side - crop_h) // 2))
            crop = padded

        sous_dossier = output_folder / nom_dossier
        sous_dossier.mkdir(parents=True, exist_ok=True)

        nom_base = f"{image_id}_{x1_crop}_{y1_crop}"
        chemin_sortie = _chemin_unique(sous_dossier, nom_base, ".jpg")
        crop.save(chemin_sortie, "JPEG", quality=95)

    return {"n_bbox": n_bbox, "n_fallback": n_fallback, "n_ignore": n_ignore, "n_trop_petit": n_trop_petit}


def _lire_cell_dims(drop_folder):
    """Retourne (cell_width, cell_height) depuis metadata.json, ou (None, None) si absent."""
    meta_path = Path(drop_folder) / "metadata.json"
    if not meta_path.exists():
        return None, None
    try:
        with open(meta_path, encoding="utf-8") as f:
            meta = json.load(f)
        return meta.get("cell_width"), meta.get("cell_height")
    except Exception:
        return None, None


def _trouver_bbox_proche(detections_folder, frame_key, x_clic, y_clic, seuil_px=100):
    """
    Cherche dans detections_folder les fichiers {frame_key}_x1_y1_x2_y2.jpg.
    Retourne (x1, y1, x2, y2) dont le centre est le plus proche du clic
    et à distance <= seuil_px. Retourne None si aucune bbox dans le seuil.
    """
    if not detections_folder.exists():
        return None

    candidats = list(detections_folder.glob(f"{frame_key}_*.jpg"))
    if not candidats:
        return None

    regex = re.compile(rf"^{re.escape(frame_key)}_(\d+)_(\d+)_(\d+)_(\d+)\.jpg$")

    meilleure_dist = float("inf")
    meilleure_bbox = None

    for fichier in candidats:
        m = regex.match(fichier.name)
        if not m:
            continue
        x1, y1, x2, y2 = int(m.group(1)), int(m.group(2)), int(m.group(3)), int(m.group(4))
        cx = (x1 + x2) / 2.0
        cy = (y1 + y2) / 2.0
        dist = ((cx - x_clic) ** 2 + (cy - y_clic) ** 2) ** 0.5
        if dist < meilleure_dist:
            meilleure_dist = dist
            meilleure_bbox = (x1, y1, x2, y2)

    if meilleure_bbox is not None and meilleure_dist <= seuil_px:
        return meilleure_bbox
    return None


def _nettoyer_nom_dossier(nom):
    """Remplace les caractères invalides dans un nom de dossier Windows/Linux."""
    nom = re.sub(r'[\\/:*?"<>|]', "_", nom)
    nom = nom.strip(". ")
    return nom or "_inconnu"


def _chemin_unique(dossier, nom_base, extension):
    """Retourne un Path libre en ajoutant _1, _2… en cas de collision."""
    chemin = dossier / f"{nom_base}{extension}"
    if not chemin.exists():
        return chemin
    compteur = 1
    while True:
        chemin = dossier / f"{nom_base}_{compteur}{extension}"
        if not chemin.exists():
            return chemin
        compteur += 1


def draw_annotations(image_path, annotations, output_path):
    image_path = Path(image_path)
    output_path = Path(output_path)

    img = Image.open(image_path).convert("RGB")
    draw = ImageDraw.Draw(img)

    for ann in annotations:
        x = ann.get("x_image")
        y = ann.get("y_image")
        if x is None or y is None:
            continue

        try:
            x = float(x)
            y = float(y)
        except Exception:
            continue

        color = str(ann.get("species_color", "#FF0000"))
        species_name = str(ann.get("species_common_name", "Unknown"))
        count = ann.get("count", 1)
        is_uncertain = _to_bool(ann.get("is_uncertain", False))

        radius = 8

        if is_uncertain:
            draw.ellipse(
                (x - radius - 5, y - radius - 5, x + radius + 5, y + radius + 5),
                outline="#FFFF00",
                width=5,
            )

        draw.ellipse(
            (x - radius, y - radius, x + radius, y + radius),
            fill=color,
            outline="black",
        )

        label = species_name
        try:
            if int(count) > 1:
                label = f"{label} x{int(count)}"
        except Exception:
            pass

        if is_uncertain:
            label = f"? {label}"

        draw.text((x + 12, y - 10), label, fill=color)

    img.save(output_path)